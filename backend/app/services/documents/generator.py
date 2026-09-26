"""
Medpark Meeting Intelligence System - Document Generation Engine
Produces immutable, versioned PDF and DOCX Minutes of Meeting supporting Romanian diacritics and Cyrillic.
"""

from pathlib import Path
from datetime import datetime
import os
import re
from typing import Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from fpdf import FPDF
from pydantic import BaseModel
from app.core.logging import logger
from app.models.meeting import Meeting
from app.models.extraction import ActionItem, EvidenceQuote, MinutesOfMeeting
from app.models.transcript import Transcript, TranscriptSegment
from app.services.extraction.attribution_render import render_minutes
from app.storage.repository import repository

# Audit banners rendered in both document formats (Romanian, as read by the clinical reviewer)
DEGRADED_BANNER = "DRAFT NEVALIDAT — LLM LOCAL INDISPONIBIL"
NAME_REVIEW_NOTICE = "Verificare nume necesară"

# Speaker attribution legend. Printed only when a reviewer confirmed at least one printable segment;
# the notice sentence is fixed wording shared with the review UI and must not be paraphrased.
ATTRIBUTION_LEGEND_TITLE = "ATRIBUIREA VORBITORILOR"
ATTRIBUTION_NOTICE = (
    "Vorbitorii neidentificați sunt marcați ca «Vorbitor N». "
    "Identificarea vocală este propusă automat și validată de un revizor uman."
)
ATTRIBUTION_REVISION_MISMATCH = (
    "Atenție: confirmarea vorbitorilor a fost făcută pentru Rev.{confirmed} iar documentul este Rev.{current}; "
    "atribuirea trebuie reverificată."
)
ATTRIBUTION_STATE_VERB = {"confirmed": "confirmat", "corrected": "corectat"}
# The transcript keeps the English cluster label ("Speaker 2"); the Romanian document prints it as
# "Vorbitor 2". Anything that is not an anonymous label passes through untouched.
ANONYMOUS_SPEAKER_LABEL = re.compile(r"^Speaker (\d+)$")
UNASSIGNED_OWNER = "Nespecificat"

# Language line of the info table ("Limbi detectate: RO, RU, EN"), printed only when the stored
# transcript has segments: an empty transcript defaults languages_detected to ["ro"], and a
# no-speech document must not claim a detected language.
LANGUAGES_LABEL = "Limbi detectate:"
LANGUAGE_COUNTS_LABEL = "Segmente per limbă:"


def format_detected_languages(transcript: Optional[Transcript]) -> Optional[str]:
    """'RO, RU, EN' from transcript.languages_detected (uppercased, ", "-joined); None without segments."""
    if transcript is None or not transcript.segments or not transcript.languages_detected:
        return None
    return ", ".join(code.upper() for code in transcript.languages_detected)


def format_language_segment_counts(transcript: Optional[Transcript]) -> Optional[str]:
    """
    'RO: 12, RU: 5, MIXED: 1' counted directly from the segments, most frequent first (ties by code).
    A segment carrying no language code is counted under UND so the counts always sum to the segment total.
    """
    if transcript is None or not transcript.segments:
        return None
    counts: dict[str, int] = {}
    for seg in transcript.segments:
        code = (seg.language or "und").upper()
        counts[code] = counts.get(code, 0) + 1
    ordered = sorted(counts.items(), key=lambda item: (-item[1], item[0]))
    return ", ".join(f"{code}: {count}" for code, count in ordered)


def truncate_quote(quote: str, limit: int) -> str:
    """
    Shortens an evidence quote for a document cell, appending the ellipsis only when the quote
    was actually cut. Marking an untouched quote as truncated would misrepresent the evidence.
    """
    text = (quote or "").strip()
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + "..."


def localize_speaker_label(label: Optional[str]) -> str:
    """Renders the anonymous cluster label in Romanian ('Speaker 2' -> 'Vorbitor 2'); other text is unchanged."""
    if not label:
        return ""
    match = ANONYMOUS_SPEAKER_LABEL.match(label.strip())
    return f"Vorbitor {match.group(1)}" if match else label.strip()


class AttributionLegendEntry(BaseModel):
    """One confirmed person as printed in the document legend (never a suggestion or a score)."""
    name: str
    state: str
    confirmed_by: Optional[str] = None
    confirmed_at: Optional[datetime] = None
    confirmed_for_revision: Optional[int] = None

    def render(self) -> str:
        verb = ATTRIBUTION_STATE_VERB.get(self.state, "confirmat")
        line = f"{self.name} — {verb} de {self.confirmed_by or 'revizor neînregistrat'}"
        if self.confirmed_at:
            line += f" la {self.confirmed_at.strftime('%d.%m.%Y %H:%M')} UTC"
        if self.confirmed_for_revision is not None:
            line += f" (Rev.{self.confirmed_for_revision})"
        return line


def build_segment_index(transcript: Optional[Transcript]) -> dict[str, TranscriptSegment]:
    """Segment id -> segment for the evidence renderers; empty when no transcript exists yet."""
    if transcript is None:
        return {}
    return {seg.id: seg for seg in transcript.segments}


def build_attribution_legend(transcript: Optional[Transcript]) -> list[AttributionLegendEntry]:
    """
    Distinct confirmed names carried by printable segments, in order of first appearance.
    Empty unless at least one segment is printable, so a document without a human confirmation
    prints no legend at all. The cluster label is deliberately NOT mapped to the name here: the
    short turns of the same cluster stay anonymous, and a "Vorbitor 2 = Dr. X" line would let a
    reader attribute them anyway.
    """
    if transcript is None:
        return []
    entries: dict[str, AttributionLegendEntry] = {}
    for seg in transcript.segments:
        if not (seg.printable_name and seg.confirmed_display_name):
            continue
        if seg.confirmed_display_name in entries:
            continue
        entries[seg.confirmed_display_name] = AttributionLegendEntry(
            name=seg.confirmed_display_name,
            state=seg.attribution_state,
            confirmed_by=seg.confirmed_by,
            confirmed_at=seg.confirmed_at,
            confirmed_for_revision=seg.confirmed_for_revision
        )
    return list(entries.values())


def render_evidence_speaker(evidence: EvidenceQuote, segment_index: dict[str, TranscriptSegment]) -> str:
    """
    Speaker label printed next to an evidence quote.
    The cited segment's display_speaker is authoritative (a name only when printable). Without the
    segment there is nothing to verify a name against, so only an anonymous label is printed; a
    name copied into the quote (even one flagged speaker_is_confirmed) fails closed to no label.
    """
    segment = segment_index.get(evidence.segment_id)
    if segment is not None:
        return localize_speaker_label(segment.display_speaker)
    if evidence.speaker and ANONYMOUS_SPEAKER_LABEL.match(evidence.speaker.strip()):
        return localize_speaker_label(evidence.speaker)
    return ""


def render_action_owner(action: ActionItem, segment_index: dict[str, TranscriptSegment]) -> str:
    """
    Responsible person as printed in the document.
    An anonymous speaker owner prints as 'Vorbitor N'. A confirmed-speaker owner is re-checked against
    the transcript: every cited segment must still print exactly that name, otherwise the owner falls
    back to the cited anonymous label (or 'Nespecificat' when nothing can be verified).
    """
    if action.owner_source == "confirmed_speaker":
        cited = [segment_index.get(ev.segment_id) for ev in action.evidence]
        if cited and all(seg is not None and seg.printable_name and seg.display_speaker == action.owner for seg in cited):
            return action.owner
        anonymous = {seg.speaker for seg in cited if seg is not None}
        return localize_speaker_label(sorted(anonymous)[0]) if len(anonymous) == 1 else UNASSIGNED_OWNER
    return localize_speaker_label(action.owner) or UNASSIGNED_OWNER


def render_minutes_for_reader(minutes: MinutesOfMeeting, transcript: Transcript) -> MinutesOfMeeting:
    """
    The minutes as a reader sees them (API responses and the PDF/DOCX): speaker tokens in the PROSE resolve at
    cluster level through render_minutes, while every action owner and evidence speaker keeps its STORED value.
    Those are maintained per segment by the speaker decisions (a name only where every cited segment is
    confirmed/corrected AND printable); resolving them through the cluster map would print a name for a turn
    below the printable floor. Returns a copy; the stored minutes are never modified.
    """
    shown = render_minutes(minutes, transcript)
    stored_items = [*minutes.decisions, *minutes.action_items, *minutes.risks_and_questions]
    shown_items = [*shown.decisions, *shown.action_items, *shown.risks_and_questions]
    for shown_item, stored_item in zip(shown_items, stored_items):
        for shown_quote, stored_quote in zip(shown_item.evidence, stored_item.evidence):
            shown_quote.speaker = stored_quote.speaker
    for shown_action, stored_action in zip(shown.action_items, minutes.action_items):
        shown_action.owner = stored_action.owner
    return shown


def format_evidence_line(evidence: EvidenceQuote, segment_index: dict[str, TranscriptSegment], limit: int) -> str:
    """'[10s-15s] Vorbitor 2: "quote"' with the speaker part omitted when no label may be printed."""
    speaker = render_evidence_speaker(evidence, segment_index)
    prefix = f"[{int(evidence.start)}s-{int(evidence.end)}s]"
    if speaker:
        prefix += f" {speaker}"
    return f"{prefix}: \"{truncate_quote(evidence.quote, limit)}\""


class PDFReport(FPDF):
    """Custom FPDF generator with Medpark styling."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.font_family_to_use = "Helvetica"
        # Extraction provenance printed on every page footer; set by the generator per document
        self.model_version = ""
        # Core Helvetica ships bold and italic faces; a Unicode TTF only gets the styles we register
        self.bold_style = "B"
        self.italic_style = "I"
        # Check cross-platform Unicode font candidates (Windows Arial, Linux/Docker DejaVu, Liberation)
        # as (regular, bold, italic) sets so emphasis survives the switch to a Unicode family.
        font_candidates = [
            (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf"), Path("C:/Windows/Fonts/ariali.ttf")),
            (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
             Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
             Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf")),
            (Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
             Path("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
             Path("/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf")),
            (Path("/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"),
             Path("/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf"),
             Path("/usr/share/fonts/truetype/noto/NotoSans-Italic.ttf")),
            (Path("/usr/share/fonts/opentype/freefont/FreeSans.ttf"),
             Path("/usr/share/fonts/opentype/freefont/FreeSansBold.ttf"),
             Path("/usr/share/fonts/opentype/freefont/FreeSansOblique.ttf")),
        ]
        for regular_path, bold_path, italic_path in font_candidates:
            if not regular_path.exists():
                continue
            try:
                self.add_font("AppUnicode", "", str(regular_path))
            except Exception:
                continue
            self.font_family_to_use = "AppUnicode"
            self.bold_style = self._register_style(bold_path, "B")
            self.italic_style = self._register_style(italic_path, "I")
            break

    def _register_style(self, font_path: Path, style: str) -> str:
        """Registers an emphasis face, returning the usable style flag ("" when unavailable)."""
        if not font_path.exists():
            return ""
        try:
            self.add_font("AppUnicode", style, str(font_path))
            return style
        except Exception:
            return ""

    def header(self):
        is_unicode = self.font_family_to_use != "Helvetica"
        self.set_font(self.font_family_to_use, self.bold_style, 13)
        self.set_text_color(0, 51, 102)  # Hospital Navy Blue

        header_title = "SPITALUL INTERNAȚIONAL MEDPARK" if is_unicode else "SPITALUL INTERNATIONAL MEDPARK"
        self.cell(0, 8, header_title, new_x="LMARGIN", new_y="NEXT", align="L")

        self.set_font(self.font_family_to_use, self.italic_style, 9)
        self.set_text_color(100, 100, 100)
        sub_text = "Sistem Inteligent de Documentare Offline a Ședințelor" if is_unicode else "Sistem Inteligent de Documentare Offline a Sedintelor"
        self.cell(0, 5, sub_text, new_x="LMARGIN", new_y="NEXT", align="L")

        self.ln(2)
        self.set_draw_color(0, 51, 102)
        self.set_line_width(0.5)
        self.line(10, self.get_y(), 200, self.get_y())
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font(self.font_family_to_use, self.italic_style, 8)
        self.set_text_color(128, 128, 128)
        foot_text = f"Pagina {self.page_no()} | Document Confidential de Uz Intern Medpark"
        if self.model_version:
            foot_text += f" | Model extragere: {self.model_version}"
        self.cell(0, 10, foot_text, align="C")


class DocumentGenerator:
    """Generates official, evidence-linked PDF and DOCX reports."""

    def _load_transcript(self, meeting: Meeting, transcript: Optional[Transcript]) -> Optional[Transcript]:
        """
        The stored transcript is the only source of a printable name: its segments decide whether an
        evidence quote or an owner may carry a person's name, and whether the legend is printed at all.
        A missing or unreadable transcript renders everything anonymous rather than aborting the export.
        """
        if transcript is not None:
            return transcript
        try:
            return repository.get_transcript(meeting.id)
        except Exception as exc:
            logger.warning(f"Transcript of meeting {meeting.id} unavailable for document rendering, printing anonymous labels: {exc}")
            return None

    def _render_for_print(self, meeting: Meeting, minutes: MinutesOfMeeting, transcript: Optional[Transcript]) -> MinutesOfMeeting:
        """
        A rendered COPY of the minutes for the page: speaker tokens in the prose become the confirmed/labelled
        name where the transcript allows it and the anonymous localized form everywhere else; owners and evidence
        speakers keep the stored per-segment values (render_action_owner / render_evidence_speaker re-check them).
        Without a transcript every token renders anonymous. The stored minutes are never modified by document generation.
        """
        return render_minutes_for_reader(minutes, transcript or Transcript(meeting_id=meeting.id))

    def generate_all(
        self,
        meeting: Meeting,
        minutes: MinutesOfMeeting,
        pdf_out: Path,
        docx_out: Path,
        transcript: Optional[Transcript] = None
    ) -> tuple[Path, Path]:
        """Generates both PDF and DOCX files for a given meeting revision."""
        transcript = self._load_transcript(meeting, transcript)
        pdf_path = self.generate_pdf(meeting, minutes, pdf_out, transcript=transcript)
        docx_path = self.generate_docx(meeting, minutes, docx_out, transcript=transcript)
        return pdf_path, docx_path

    def generate_docx(
        self,
        meeting: Meeting,
        minutes: MinutesOfMeeting,
        output_path: Path,
        transcript: Optional[Transcript] = None
    ) -> Path:
        """Constructs formatted Word document (.docx) with tables and metadata."""
        output_path.parent.mkdir(parents=True, exist_ok=True)
        transcript = self._load_transcript(meeting, transcript)
        minutes = self._render_for_print(meeting, minutes, transcript)
        segment_index = build_segment_index(transcript)
        legend = build_attribution_legend(transcript)
        doc = Document()

        # Document Header
        title_p = doc.add_paragraph()
        title_run = title_p.add_run("PROCES-VERBAL AL ȘEDINȚEI (MINUTES OF MEETING)")
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_run.font.color.rgb = RGBColor(0, 51, 102)

        subtitle_p = doc.add_paragraph()
        sub_run = subtitle_p.add_run(f"Spitalul Internațional Medpark | {meeting.meeting_type.value.upper()}")
        sub_run.font.size = Pt(11)
        sub_run.font.color.rgb = RGBColor(100, 100, 100)

        # Audit banners: a heuristic draft and unconfirmed names must be the first thing a reader sees
        if minutes.is_degraded:
            warn_p = doc.add_paragraph()
            warn_run = warn_p.add_run(DEGRADED_BANNER)
            warn_run.bold = True
            warn_run.font.color.rgb = RGBColor(192, 0, 0)
        if minutes.needs_name_review:
            review_p = doc.add_paragraph()
            review_run = review_p.add_run(NAME_REVIEW_NOTICE)
            review_run.bold = True
            review_run.font.color.rgb = RGBColor(184, 134, 11)

        # Meeting Info Grid
        doc.add_heading("Informații Generale", level=2)
        info_data = [
            ("Titlu Ședință:", meeting.title),
            ("Participanți:", ", ".join([f"{a.name} ({a.department})" if a.department else a.name for a in meeting.attendees]) if meeting.attendees else "Conform foii de prezență"),
            ("Model extragere:", minutes.model_version)
        ]
        # Languages of the stored transcript (per-window ASR language identification), omitted without segments
        languages_line = format_detected_languages(transcript)
        if languages_line:
            info_data.append((LANGUAGES_LABEL, languages_line))
            info_data.append((LANGUAGE_COUNTS_LABEL, format_language_segment_counts(transcript) or ""))
        info_table = doc.add_table(rows=len(info_data), cols=2)
        info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row_idx, (label, val) in enumerate(info_data):
            info_table.rows[row_idx].cells[0].paragraphs[0].add_run(label).bold = True
            info_table.rows[row_idx].cells[1].paragraphs[0].add_run(val)

        # Executive Summary
        doc.add_heading("Rezumat Executiv", level=2)
        doc.add_paragraph(minutes.summary_ro)
        if minutes.summary_ru:
            p_ru = doc.add_paragraph()
            r_ru = p_ru.add_run(f"[RU Резюме] {minutes.summary_ru}")
            r_ru.italic = True
        if minutes.summary_en:
            p_en = doc.add_paragraph()
            r_en = p_en.add_run(f"[EN Summary] {minutes.summary_en}")
            r_en.italic = True

        # Decisions Section
        doc.add_heading("Decizii Adoptate", level=2)
        if minutes.decisions:
            dec_table = doc.add_table(rows=1, cols=3)
            dec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = dec_table.rows[0].cells
            hdr[0].paragraphs[0].add_run("Domeniu").bold = True
            hdr[1].paragraphs[0].add_run("Decizie Adoptată").bold = True
            hdr[2].paragraphs[0].add_run("Dovadă Audio (Timestamp)").bold = True

            for d in minutes.decisions:
                row = dec_table.add_row().cells
                row[0].paragraphs[0].add_run(d.topic)
                row[1].paragraphs[0].add_run(d.decision)
                ev_text = "; ".join([format_evidence_line(e, segment_index, 40) for e in d.evidence]) if d.evidence else "N/A"
                row[2].paragraphs[0].add_run(ev_text)
        else:
            doc.add_paragraph("Nicio decizie formală înregistrată.")

        # Action Items Section
        doc.add_heading("Plan de Acțiuni & Sarcini (Action Items)", level=2)
        if minutes.action_items:
            act_table = doc.add_table(rows=1, cols=4)
            act_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = act_table.rows[0].cells
            hdr[0].paragraphs[0].add_run("Sarcină / Activitate").bold = True
            hdr[1].paragraphs[0].add_run("Responsabil").bold = True
            hdr[2].paragraphs[0].add_run("Termen Limită").bold = True
            hdr[3].paragraphs[0].add_run("Prioritate").bold = True

            for a in minutes.action_items:
                row = act_table.add_row().cells
                row[0].paragraphs[0].add_run(a.task)
                row[1].paragraphs[0].add_run(render_action_owner(a, segment_index))
                deadline_str = a.deadline_date or a.deadline_phrase or "Nespecificat"
                row[2].paragraphs[0].add_run(deadline_str)
                row[3].paragraphs[0].add_run(a.priority.upper())
        else:
            doc.add_paragraph("Nicio sarcină de lucru identificată.")

        # Speaker attribution legend: only confirmed names, only when at least one segment prints one
        if legend:
            doc.add_heading("Atribuirea Vorbitorilor", level=2)
            for entry in legend:
                doc.add_paragraph(entry.render(), style="List Bullet")
            mismatched = [e for e in legend if e.confirmed_for_revision is not None and e.confirmed_for_revision != minutes.revision]
            if mismatched:
                warn_p = doc.add_paragraph()
                warn_run = warn_p.add_run(ATTRIBUTION_REVISION_MISMATCH.format(
                    confirmed=", ".join(str(e.confirmed_for_revision) for e in mismatched), current=minutes.revision
                ))
                warn_run.bold = True
                warn_run.font.color.rgb = RGBColor(184, 134, 11)
            notice_p = doc.add_paragraph()
            notice_p.add_run(ATTRIBUTION_NOTICE).italic = True

        doc.save(str(output_path))
        logger.info(f"Generated DOCX minutes: {output_path}")
        return output_path

    def generate_pdf(
        self,
        meeting: Meeting,
        minutes: MinutesOfMeeting,
        output_path: Path,
        transcript: Optional[Transcript] = None
    ) -> Path:
        """Constructs official PDF document using FPDF."""
        output_path.parent.mkdir(parents=True, exist_ok=True)
        transcript = self._load_transcript(meeting, transcript)
        minutes = self._render_for_print(meeting, minutes, transcript)
        segment_index = build_segment_index(transcript)
        legend = build_attribution_legend(transcript)
        pdf = PDFReport()
        pdf.model_version = minutes.model_version
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)

        use_font = pdf.font_family_to_use
        is_unicode = use_font != "Helvetica"
        style_b = pdf.bold_style
        style_i = pdf.italic_style

        def safe_text(txt: str) -> str:
            return txt if is_unicode else self._sanitize_text(txt)

        # Title. Variable-length content always uses multi_cell: cell() clips at the page width
        # instead of wrapping, which silently truncates long meeting titles and decisions.
        pdf.set_font(use_font, style_b, 13)
        pdf.set_x(10)
        pdf.multi_cell(190, 8, safe_text(f"PROCES-VERBAL: {meeting.title.upper()}"), new_x="LMARGIN", new_y="NEXT", align="L")

        # Audit banners directly under the title (red draft warning, amber name-review notice)
        if minutes.is_degraded:
            pdf.set_font(use_font, style_b, 10)
            pdf.set_text_color(192, 0, 0)
            pdf.set_x(10)
            pdf.multi_cell(190, 6, safe_text(DEGRADED_BANNER), new_x="LMARGIN", new_y="NEXT", align="L")
            pdf.set_text_color(0, 0, 0)
        if minutes.needs_name_review:
            pdf.set_font(use_font, style_b, 9)
            pdf.set_text_color(184, 134, 11)
            pdf.set_x(10)
            pdf.multi_cell(190, 5, safe_text(NAME_REVIEW_NOTICE), new_x="LMARGIN", new_y="NEXT", align="L")
            pdf.set_text_color(0, 0, 0)

        pdf.set_font(use_font, "", 10)
        pdf.cell(0, 6, safe_text(f"Data: {meeting.scheduled_at.strftime('%Y-%m-%d %H:%M')} | Tip: {meeting.meeting_type.value.upper()} | Rev.{minutes.revision}"), new_x="LMARGIN", new_y="NEXT")
        if meeting.attendees:
            att_str = ", ".join(f"{a.name} ({a.department})" if a.department else a.name for a in meeting.attendees)
            pdf.set_x(10)
            pdf.multi_cell(190, 5, safe_text(f"Participanți: {att_str}"), new_x="LMARGIN", new_y="NEXT")
        # Languages of the stored transcript (per-window ASR language identification), omitted without segments
        languages_line = format_detected_languages(transcript)
        if languages_line:
            pdf.set_x(10)
            pdf.multi_cell(
                190, 6,
                safe_text(f"{LANGUAGES_LABEL} {languages_line} | {LANGUAGE_COUNTS_LABEL} {format_language_segment_counts(transcript)}"),
                new_x="LMARGIN", new_y="NEXT"
            )
        pdf.ln(3)

        # Executive Summary Section
        pdf.set_font(use_font, style_b, 11)
        pdf.set_fill_color(240, 244, 248)
        pdf.cell(190, 7, safe_text("REZUMAT EXECUTIV / EXECUTIVE SUMMARY"), fill=True, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(use_font, "", 9)
        pdf.set_x(10)
        pdf.multi_cell(190, 5, safe_text(minutes.summary_ro), new_x="LMARGIN", new_y="NEXT")
        if minutes.summary_ru:
            pdf.ln(1)
            pdf.set_font(use_font, style_i, 8)
            pdf.set_text_color(60, 60, 60)
            pdf.set_x(10)
            pdf.multi_cell(190, 4.5, safe_text(f"[RU] {minutes.summary_ru}"), new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(0, 0, 0)
        if minutes.summary_en:
            pdf.ln(1)
            pdf.set_font(use_font, style_i, 8)
            pdf.set_text_color(60, 60, 60)
            pdf.set_x(10)
            pdf.multi_cell(190, 4.5, safe_text(f"[EN] {minutes.summary_en}"), new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(0, 0, 0)
        pdf.ln(3)

        # Decisions
        pdf.set_font(use_font, style_b, 11)
        pdf.set_fill_color(240, 244, 248)
        pdf.cell(190, 7, safe_text(f"DECIZII ADOPTATE ({len(minutes.decisions)})"), fill=True, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(use_font, "", 9)
        if minutes.decisions:
            for i, d in enumerate(minutes.decisions, 1):
                pdf.set_font(use_font, style_b, 9)
                pdf.set_x(10)
                pdf.multi_cell(190, 5, safe_text(f"{i}. [{d.topic}]"), new_x="LMARGIN", new_y="NEXT")
                pdf.set_font(use_font, "", 9)
                pdf.set_x(10)
                pdf.multi_cell(190, 5, safe_text(f"   {d.decision}"), new_x="LMARGIN", new_y="NEXT")
                if d.evidence:
                    pdf.set_font(use_font, style_i, 8)
                    pdf.set_text_color(90, 90, 90)
                    ev_str = f"   Audio Dovada {format_evidence_line(d.evidence[0], segment_index, 70)}"
                    pdf.set_x(10)
                    pdf.multi_cell(190, 4, safe_text(ev_str), new_x="LMARGIN", new_y="NEXT")
                    pdf.set_text_color(0, 0, 0)
                pdf.ln(2)
        else:
            pdf.cell(190, 5, safe_text("Nicio decizie formala."), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

        # Action Items
        pdf.set_font(use_font, style_b, 11)
        pdf.set_fill_color(240, 244, 248)
        pdf.cell(190, 7, safe_text(f"PLAN DE ACTIUNI & SARCINI ({len(minutes.action_items)})"), fill=True, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(use_font, "", 9)
        if minutes.action_items:
            for i, a in enumerate(minutes.action_items, 1):
                pdf.set_font(use_font, style_b, 9)
                deadline_info = f" | Termen: {a.deadline_date or a.deadline_phrase or 'Nespecificat'}"
                pdf.set_x(10)
                pdf.multi_cell(190, 5, safe_text(f"{i}. Responsabil: {render_action_owner(a, segment_index)}{deadline_info} [{a.priority.upper()}]"), new_x="LMARGIN", new_y="NEXT")
                pdf.set_font(use_font, "", 9)
                pdf.set_x(10)
                pdf.multi_cell(190, 5, safe_text(f"   Sarcina: {a.task}"), new_x="LMARGIN", new_y="NEXT")
                if a.evidence:
                    pdf.set_font(use_font, style_i, 8)
                    pdf.set_text_color(90, 90, 90)
                    ev_str = f"   Audio Dovada {format_evidence_line(a.evidence[0], segment_index, 70)}"
                    pdf.set_x(10)
                    pdf.multi_cell(190, 4, safe_text(ev_str), new_x="LMARGIN", new_y="NEXT")
                    pdf.set_text_color(0, 0, 0)
                pdf.ln(2)
        else:
            pdf.cell(190, 5, safe_text("Nicio sarcina de lucru."), new_x="LMARGIN", new_y="NEXT")

        # Speaker attribution legend: only confirmed names, only when at least one segment prints one
        if legend:
            pdf.ln(3)
            pdf.set_font(use_font, style_b, 11)
            pdf.set_fill_color(240, 244, 248)
            pdf.cell(190, 7, safe_text(ATTRIBUTION_LEGEND_TITLE), fill=True, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font(use_font, "", 9)
            for entry in legend:
                pdf.set_x(10)
                pdf.multi_cell(190, 5, safe_text(f"- {entry.render()}"), new_x="LMARGIN", new_y="NEXT")
            mismatched = [e for e in legend if e.confirmed_for_revision is not None and e.confirmed_for_revision != minutes.revision]
            if mismatched:
                pdf.set_font(use_font, style_b, 9)
                pdf.set_text_color(184, 134, 11)
                pdf.set_x(10)
                pdf.multi_cell(190, 5, safe_text(ATTRIBUTION_REVISION_MISMATCH.format(
                    confirmed=", ".join(str(e.confirmed_for_revision) for e in mismatched), current=minutes.revision
                )), new_x="LMARGIN", new_y="NEXT")
                pdf.set_text_color(0, 0, 0)
            pdf.set_font(use_font, style_i, 8)
            pdf.set_text_color(90, 90, 90)
            pdf.set_x(10)
            pdf.multi_cell(190, 4, safe_text(ATTRIBUTION_NOTICE), new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(0, 0, 0)

        pdf.output(str(output_path))
        logger.info(f"Generated PDF minutes: {output_path}")
        return output_path

    def _sanitize_text(self, text: str) -> str:
        """Replaces characters incompatible with standard latin-1 Helvetica while preserving readability."""
        replacements = {
            "ș": "s", "Ș": "S",
            "ț": "t", "Ț": "T",
            "ă": "a", "Ă": "A",
            "î": "i", "Î": "I",
            "â": "a", "Â": "A",
            "—": "-", "–": "-",
            "„": '"', "”": '"',
            "«": '"', "»": '"',
            "’": "'", "‘": "'"
        }
        res = text
        for k, v in replacements.items():
            res = res.replace(k, v)
        return res.encode("latin-1", "replace").decode("latin-1")


document_generator = DocumentGenerator()
